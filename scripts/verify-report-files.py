"""Read-only cross-format verification of Rust-generated synthetic report fixtures."""
import csv
import json
import re
import sys
from decimal import Decimal
from pathlib import Path
from pypdf import PdfReader
from openpyxl import load_workbook
from docx import Document

def raw(cell):
    if cell['kind'] == 'text':
        return cell['value']
    scale = 2 if cell['kind'] == 'money' else 4 if cell['kind'] == 'percentage' else cell['scale']
    return format(Decimal(cell['value']).scaleb(-scale), f'.{scale}f')

def display(cell):
    value = raw(cell)
    if cell['kind'] == 'percentage':
        return value + '%'
    if cell['kind'] == 'money':
        return ('-' if value.startswith('-') else '') + '$' + format(abs(Decimal(value)), ',.2f')
    return value

def compact(text):
    return re.sub(r'\s+', '', text)

root = Path(sys.argv[1])
for path in sorted(root.glob('*.json')):
    if len(sys.argv)>2 and path.stem != sys.argv[2]:
        continue
    report = json.loads(path.read_text(encoding='utf-8-sig'))
    if not isinstance(report, dict) or 'sections' not in report:
        continue
    with path.with_suffix('.csv').open(encoding='utf-8', newline='') as stream:
        rows = list(csv.DictReader(stream))
    expected = sum(len(s['rows']) for s in report['sections'])
    assert len(rows) == expected, path
    index = 0
    for section in report['sections']:
        for row in section['rows']:
            assert rows[index]['Section'] == section['name']
            for label, cell in zip(section['columns'], row):
                text = raw(cell)
                if cell['kind'] == 'text' and text.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')):
                    text = "'" + text
                assert rows[index][label] == text, (path, label, rows[index][label], text)
            index += 1
    workbook = load_workbook(path.with_suffix('.xlsx'), data_only=False)
    for section in report['sections']:
        sheet = workbook['Settlement' if report['kind'] == 'settlement_detail' else section['name']]
        cells = list(sheet.values)
        start = next(i for i, row in enumerate(cells) if i >= 4 and row[0] == section['name']) + 2
        assert list(cells[start-1][:len(section['columns'])]) == section['columns']
        for i, row in enumerate(section['rows']):
            for j, cell in enumerate(row):
                actual = cells[start+i][j]
                assert sheet.cell(start+i+1, j+1).data_type != 'f'
                if isinstance(actual, (int, float)):
                    assert Decimal(str(actual)) == Decimal(raw(cell)), (path, actual, cell)
                else:
                    assert (actual or '') == display(cell), (path, actual, cell)
    document = Document(path.with_suffix('.docx'))
    assert len(document.tables) == len(report['sections'])
    for table, section in zip(document.tables, report['sections']):
        assert [c.text for c in table.rows[0].cells] == section['columns']
        for row, expected_row in zip(table.rows[1:], section['rows']):
            assert [c.text for c in row.cells] == [display(c) for c in expected_row], path
    pdf = PdfReader(path.with_suffix('.pdf'))
    text = '\n'.join(p.extract_text() for p in pdf.pages)
    assert compact(report['title']) in compact(text)
    for section in report['sections']:
        for row in section['rows']:
            for cell in row:
                assert compact(display(cell)) in compact(text), (path, display(cell))
    if report['draft']:
        assert all('DRAFT / NOT FINALIZED' in p.extract_text() for p in pdf.pages)
    if report['kind'] == 'owner_payroll':
        assert 'ACH Charge' not in text
        assert not any('ACH' in column for column in rows[0])
    print(f'PASS {path.stem}: {expected} rows agree across CSV/XLSX/DOCX/PDF; {len(pdf.pages)} PDF pages')
